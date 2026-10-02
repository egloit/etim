"""
Problem list ("Fehlerliste") for a generated export file - stored with the
file (etim10_export_files.issues / gs1_export_files.issues) and offered as
an Excel download, so the user can see per article what went wrong and
clarify it with data management.
"""
import io
import re
from collections import Counter
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

# "<matnr>: <text>" (ETIM10) or "Artikel <matnr>: <text>" (GS1)
_ARTICLE = re.compile(r"^(?:Artikel )?(?P<matnr>[^\s:]+): (?P<text>.+)$", re.S)
_FEATURE = re.compile(r"\b(EF\d{6}|ZZTYPEN|EAN11|MAKTX|DAM AssetCode \w+|Zolltarifnummer|Ursprungsland)\b")
_VALUE = re.compile(r"'([^']*)'")


def collect_issues(warnings: list[str], validation_errors: Optional[list[dict]] = None) -> list[dict]:
    """Warnings + XSD errors as rows {matnr, feature, value, message, source}."""
    rows = []
    for warning in warnings:
        m = _ARTICLE.match(warning)
        matnr, text = (m["matnr"], m["text"]) if m else ("", warning)
        feature = _FEATURE.search(text)
        value = _VALUE.search(text)
        rows.append({"matnr": matnr, "feature": feature.group(1) if feature else "",
                     "value": value.group(1) if value else "", "message": text, "source": "Warnung"})
    for error in validation_errors or []:
        rows.append({"matnr": error.get("matnr") or "", "feature": error.get("element") or "",
                     "value": "", "message": error.get("message") or "", "source": "XSD-Fehler"})
    return rows


def _sheet(ws, header: list[str], rows: list[list], widths: list[int]) -> None:
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append(row)
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions


def to_xlsx(issues: list[dict], title: str) -> bytes:
    """Excel with sheet "Probleme" (one row per article and problem) and
    "Übersicht" (problems grouped by message, with number of articles)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Probleme"
    rows = sorted(issues, key=lambda r: (r["matnr"] == "", r["matnr"], r["feature"]))
    _sheet(ws, ["Materialnummer", "Merkmal/Feld", "Wert", "Meldung", "Art"],
           [[r["matnr"], r["feature"], r["value"], r["message"], r["source"]] for r in rows],
           [18, 22, 18, 110, 12])

    counts = Counter((r["source"], r["message"]) for r in issues)
    articles: dict[tuple, list[str]] = {}
    for r in issues:
        if r["matnr"]:
            articles.setdefault((r["source"], r["message"]), []).append(r["matnr"])
    overview = [[n, len(articles.get(key, [])), key[0], key[1], ", ".join(articles.get(key, [])[:20])]
                for key, n in counts.most_common()]
    _sheet(wb.create_sheet("Übersicht"), ["Anzahl", "Artikel", "Art", "Meldung", "Beispiele (max. 20)"],
           overview, [10, 10, 12, 110, 60])
    wb.properties.title = title

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
