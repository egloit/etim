import io

from openpyxl import load_workbook

from issue_report import collect_issues, to_xlsx


def test_collect_issues_parses_etim_and_gs1_warnings():
    rows = collect_issues(
        ["30658: EF004284 (Material cover): Wert 'EV001752' ist für EC001743 nicht zulässig – ausgelassen.",
         "Artikel 43706: keine BrickId-Zuordnung für ZZTYPEN=ZZTYPEN_LIK.",
         "902162: Zolltarifnummer in SAP (MARC, Werk 0090) leer.",
         "3 Artikel ohne Preis – Preisblock nur mit Datum."],
        [{"matnr": "69584", "element": "/BMECAT/PRODUCT", "message": "Element 'X': not expected."}],
    )
    assert [(r["matnr"], r["feature"], r["value"], r["source"]) for r in rows] == [
        ("30658", "EF004284", "EV001752", "Warnung"),
        ("43706", "ZZTYPEN", "", "Warnung"),
        ("902162", "Zolltarifnummer", "", "Warnung"),
        ("", "", "", "Warnung"),
        ("69584", "/BMECAT/PRODUCT", "", "XSD-Fehler"),
    ]


def test_to_xlsx_has_detail_and_overview_sheets():
    issues = collect_issues(["1: EF004282 (Light outlet): Wert 'EV003775' ist für EC001744 nicht zulässig.",
                             "2: EF004282 (Light outlet): Wert 'EV003775' ist für EC001744 nicht zulässig.",
                             "2: Zolltarifnummer in SAP (MARC, Werk 0090) leer."])
    wb = load_workbook(io.BytesIO(to_xlsx(issues, "test")))
    assert wb.sheetnames == ["Probleme", "Übersicht"]
    detail = list(wb["Probleme"].values)
    assert detail[0] == ("Materialnummer", "Merkmal/Feld", "Wert", "Meldung", "Art")
    assert [r[0] for r in detail[1:]] == ["1", "2", "2"]
    overview = list(wb["Übersicht"].values)
    assert overview[1][:2] == (2, 2) and "EF004282" in overview[1][3]
